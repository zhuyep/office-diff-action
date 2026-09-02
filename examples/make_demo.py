"""Create synthetic, redistribution-safe files for the README demo."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches as DocxInches
from docx.shared import Pt as DocxPt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent / "generated"


def add_text(slide, text, x, y, width, height, size, color, bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = PP_ALIGN.LEFT
    run = paragraph.runs[0]
    run.font.name = "DejaVu Sans"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)
    return box


def add_card(slide, title, value, x, accent):
    shape = slide.shapes.add_shape(1, Inches(x), Inches(3.25), Inches(3.55), Inches(2.05))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(248, 250, 252)
    shape.line.color.rgb = RGBColor(226, 232, 240)
    add_text(slide, title, x + 0.25, 3.52, 3.0, 0.35, 13, (71, 85, 105))
    add_text(slide, value, x + 0.25, 4.05, 2.9, 0.65, 30, accent, bold=True)


def make_deck(path: Path, current: bool) -> None:
    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = RGBColor(255, 255, 255)

    add_text(slide, "Release readiness", 0.7, 0.65, 8.0, 0.7, 30, (15, 23, 42), bold=True)
    subtitle = (
        "The document now renders, passes review, and is ready to ship."
        if current
        else "The document renders and is ready for review."
    )
    add_text(slide, subtitle, 0.72, 1.55, 8.7, 0.55, 17, (71, 85, 105))

    badge = slide.shapes.add_shape(5, Inches(10.35), Inches(0.7), Inches(2.25), Inches(0.55))
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(220, 252, 231) if current else RGBColor(254, 243, 199)
    badge.line.fill.background()
    add_text(
        slide,
        "READY" if current else "IN REVIEW",
        10.72 if current else 10.56,
        0.82,
        1.65,
        0.3,
        13,
        (22, 101, 52) if current else (146, 64, 14),
        bold=True,
    )

    add_card(slide, "Pages checked", "24" if current else "18", 0.72, (37, 99, 235))
    status_color = (22, 163, 74) if current else (220, 38, 38)
    add_card(slide, "Layout errors", "0" if current else "3", 4.89, status_color)
    add_card(slide, "Review status", "PASS" if current else "OPEN", 9.06, (124, 58, 237))

    add_text(
        slide,
        "Office Diff catches visual changes that a normal Git diff cannot show.",
        0.72,
        6.55,
        11.5,
        0.4,
        14,
        (100, 116, 139),
    )
    deck.core_properties.title = "Office Diff synthetic demo"
    deck.core_properties.subject = "Safe public test fixture"
    deck.core_properties.author = "Office Diff contributors"
    deck.core_properties.last_modified_by = "Office Diff contributors"
    deck.save(path)


def make_document(path: Path, current: bool) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = DocxInches(0.8)
    section.bottom_margin = DocxInches(0.8)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Release Readiness Report")
    title_run.bold = True
    title_run.font.size = DocxPt(22)
    intro = document.add_paragraph(
        "This synthetic file demonstrates how Office Diff reviews a Word document."
    )
    intro.paragraph_format.space_after = DocxPt(12)
    table = document.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    values = [
        ("Check", "Result"),
        ("Pages reviewed", "24" if current else "18"),
        ("Layout errors", "0" if current else "3"),
        ("Status", "PASS" if current else "OPEN"),
    ]
    for row, pair in zip(table.rows, values):
        row.cells[0].text = pair[0]
        row.cells[1].text = pair[1]
    conclusion = document.add_paragraph()
    conclusion.paragraph_format.space_before = DocxPt(12)
    conclusion.add_run(
        "The document is ready to ship." if current else "The document still needs review."
    )
    document.core_properties.title = "Office Diff synthetic demo"
    document.core_properties.author = "Office Diff contributors"
    document.core_properties.last_modified_by = "Office Diff contributors"
    document.save(path)


def main() -> None:
    ROOT.mkdir(parents=True, exist_ok=True)
    make_deck(ROOT / "before.pptx", current=False)
    make_deck(ROOT / "after.pptx", current=True)
    make_document(ROOT / "before.docx", current=False)
    make_document(ROOT / "after.docx", current=True)
    print(ROOT / "before.pptx")
    print(ROOT / "after.pptx")
    print(ROOT / "before.docx")
    print(ROOT / "after.docx")


if __name__ == "__main__":
    main()
